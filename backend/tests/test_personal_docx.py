from __future__ import annotations

import os
import subprocess
import sys
import textwrap


def test_docx_import_preserves_blocks_original_and_cited_table_row(tmp_path) -> None:
    env = os.environ.copy()
    env.update(
        {
            "LMWK_DATA_DIR": str(tmp_path / "word-library"),
            "LMWK_FIXTURE_DIR": str(tmp_path / "missing-fixture"),
            "LMWK_TESTING": "true",
            "LMWK_SEED_FIXTURE": "false",
            "DEEPSEEK_API_KEY": "",
        }
    )
    script = textwrap.dedent(
        """
        import hashlib
        import io
        import time
        from zipfile import ZIP_DEFLATED, ZipFile
        from docx import Document
        from fastapi.testclient import TestClient
        from app.adapters.deepseek import CompilePayload
        from app.main import app
        from app.services import compilation
        from app.services.settings import current_profile

        class OfflineClient:
            async def compile_knowledge(self, *, evidence, existing_knowledge):
                assert not existing_knowledge
                row = next(item for item in evidence if 'Python files must close' in item['quote'])
                assert row['topic'] == 'Python'
                assert row['blockLabel'] == '表格 1 · 第 2 行'
                assert row['blockNumber'] == 4
                return CompilePayload.model_validate({
                    'knowledge_items': [{
                        'title': 'Python file cleanup', 'type': 'rule', 'domain': 'Python',
                        'summary': 'Python files must close after a with block to avoid resource leaks.',
                        'markdown': '# Python file cleanup\\n\\nClose Python files after a with block.',
                        'evidence_ids': [row['evidenceId']],
                    }],
                    'relations': [],
                })

        def offline_client(session, *, require_available):
            return current_profile(session), OfflineClient()
        compilation.get_client = offline_client

        document = Document()
        document.add_heading('Python file lifecycle', level=1)
        document.add_paragraph('Context managers release resources at the end of a with block, even when exceptions are raised.')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Rule'
        table.cell(0, 1).text = 'Meaning'
        table.cell(1, 0).text = 'Python files must close after a with block'
        table.cell(1, 1).text = 'The operating system can reuse file handles safely.'
        document.add_paragraph('Additional paragraphs remain after the table in document order for citation review.')
        buffer = io.BytesIO()
        document.save(buffer)
        original = buffer.getvalue()
        with TestClient(app) as client:
            too_many_entries = io.BytesIO()
            with ZipFile(too_many_entries, 'w', compression=ZIP_DEFLATED) as archive:
                for index in range(2001):
                    archive.writestr(f'word/extra-{index}.xml', 'x')
            oversized_package = client.post('/api/v1/sources/import', data={'topic': 'Python'},
                files=[('files', ('too-many.docx', too_many_entries.getvalue(),
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))])
            assert oversized_package.json()['items'][0]['message'] == 'Word 文件展开后超过处理限制'
            invalid = client.post('/api/v1/sources/import', data={'topic': 'Python'},
                files=[('files', ('broken.docx', b'not an OOXML package',
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))])
            assert invalid.status_code == 202
            assert invalid.json()['acceptedCount'] == 0
            assert invalid.json()['items'][0]['status'] == 'failed'
            assert client.get('/api/v1/bootstrap').json()['project']['sourceCount'] == 0

            uploaded = client.post('/api/v1/sources/import', data={'topic': 'Python'},
                files=[('files', ('python-guidance.docx', original,
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))])
            assert uploaded.status_code == 202, uploaded.text
            payload = uploaded.json()
            assert payload['acceptedCount'] == 1
            source_id = payload['items'][0]['sourceId']
            version_id = payload['items'][0]['sourceVersionId']
            content = client.get(f'/api/v1/source-versions/{version_id}/content').json()
            assert content['sha256'] == hashlib.sha256(original).hexdigest()
            assert content['pageSpans'] == []
            assert [block['label'] for block in content['blockSpans']] == [
                '第 1 段', '第 2 段', '表格 1 · 第 1 行', '表格 1 · 第 2 行', '第 3 段',
            ]
            assert all(content['content'][block['charStart']:block['charEnd']].strip()
                       for block in content['blockSpans'])
            original_response = client.get(f'/api/v1/source-versions/{version_id}/original')
            assert original_response.status_code == 200
            assert original_response.content == original
            assert 'attachment' in original_response.headers['content-disposition']

            run_id = payload['runId']
            deadline = time.monotonic() + 8
            while time.monotonic() < deadline:
                run = client.get(f'/api/v1/compile-runs/{run_id}').json()
                if run['status'] in {'awaiting_review', 'failed'}:
                    break
                time.sleep(0.05)
            assert run['status'] == 'awaiting_review', run
            published = client.post(f'/api/v1/compile-runs/{run_id}/accept')
            assert published.status_code == 200, published.text
            knowledge_id = client.get('/api/v1/knowledge').json()['items'][0]['knowledgeId']
            detail = client.get(f'/api/v1/knowledge/{knowledge_id}').json()
            evidence = detail['evidence'][0]
            assert evidence['sourceId'] == source_id
            assert evidence['sourceVersionId'] == version_id
            assert evidence['pageNumber'] is None
            assert evidence['blockLabel'] == '表格 1 · 第 2 行'
            assert evidence['blockNumber'] == 4
            evidence_detail = client.get(f"/api/v1/evidence/{evidence['evidenceId']}").json()
            assert evidence_detail['accessible'] is True
            assert evidence['quote'] in evidence_detail['context']
        """
    )
    result = subprocess.run(
        [sys.executable, "-c", script], env=env, capture_output=True,
        text=True, timeout=30, check=False,
    )
    assert result.returncode == 0, (result.stdout + result.stderr)[-3000:]
